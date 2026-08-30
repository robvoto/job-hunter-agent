"""Build the local O*NET occupation/title taxonomy used by the title gate.

Supported inputs:
- Current/full O*NET database JSON zip (preferred).
- Current/full O*NET database text zip, including nested release folders.
- Legacy OccupationalListings.zip for backwards-compatible manual imports.

Runtime classification remains fully local. Network freshness is handled separately by
``job_hunter_agent.onet_taxonomy_refresh``.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import re
import zipfile
from pathlib import Path, PurePosixPath
from tempfile import TemporaryDirectory
from typing import Any

from openpyxl import load_workbook

from job_hunter_agent.paths import REPO_ROOT

TAXONOMY_VERSION = "O*NET-SOC 2019"
DEFAULT_OUTPUT_DIR = REPO_ROOT / "data" / "knowledge" / "occupation_taxonomy"

_LISTINGS_OCCUPATIONS_MEMBER = "2019_Occupations.xlsx"
_LISTINGS_ALT_TITLES_MEMBER = "2019_Alt_Titles.xlsx"
_FULLDB_TEXT_OCCUPATIONS_MEMBER = "Occupation Data.txt"
_FULLDB_TEXT_JOB_TITLES_MEMBER = "Job Titles.txt"
_FULLDB_TEXT_ALT_TITLES_MEMBER = "Alternate Titles.txt"
_FULLDB_TEXT_REPORTED_TITLES_MEMBER = "Sample of Reported Titles.txt"
_FULLDB_JSON_OCCUPATIONS_MEMBER = "occupation_data.json"
_FULLDB_JSON_JOB_TITLES_MEMBER = "job_titles.json"
_FULLDB_JSON_REPORTED_TITLES_MEMBER = "sample_of_reported_titles.json"
_RELEASE_RE = re.compile(r"db_(\d+)_(\d+)(?:_|\.)", re.IGNORECASE)


def normalize_title(value: str | None) -> str:
    """Normalize a role title for deterministic taxonomy lookup keys."""
    text = str(value or "").lower().replace("&", " and ")
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def _usable_title(value: str | None) -> str:
    text = str(value or "").strip()
    return "" if text.lower() in {"", "n/a", "na", "none", "null"} else text


def _member_by_basename(archive: zipfile.ZipFile, basename: str) -> str | None:
    matches = [name for name in archive.namelist() if PurePosixPath(name).name == basename]
    if not matches:
        return None
    if len(matches) > 1:
        raise ValueError(f"O*NET archive contains multiple '{basename}' members: {matches}")
    return matches[0]


def _require_member(archive: zipfile.ZipFile, basename: str) -> str:
    member = _member_by_basename(archive, basename)
    if member is None:
        raise ValueError(f"O*NET archive is missing required member '{basename}'")
    return member


def _infer_database_release(zip_path: Path) -> str | None:
    match = _RELEASE_RE.search(zip_path.name)
    if not match:
        return None
    return f"{int(match.group(1))}.{int(match.group(2))}"


def _read_xlsx_from_zip(zip_path: Path, basename: str) -> list[dict[str, Any]]:
    with zipfile.ZipFile(zip_path) as archive, TemporaryDirectory() as tmp_dir:
        member = _require_member(archive, basename)
        extracted = Path(tmp_dir) / basename
        extracted.write_bytes(archive.read(member))
        workbook = load_workbook(extracted, read_only=True, data_only=True)
        try:
            sheet = workbook[workbook.sheetnames[0]]
            rows = list(sheet.iter_rows(min_row=4, values_only=True))
        finally:
            workbook.close()
    if not rows:
        return []
    headers = [str(cell).strip() for cell in rows[0]]
    return [dict(zip(headers, row)) for row in rows[1:] if row and row[0]]


def _read_txt_table(archive: zipfile.ZipFile, basename: str) -> list[dict[str, str]]:
    member = _require_member(archive, basename)
    raw = archive.read(member).decode("utf-8-sig")
    return [row for row in csv.DictReader(io.StringIO(raw), delimiter="\t") if row]


def _read_json_table(archive: zipfile.ZipFile, basename: str) -> list[dict[str, Any]]:
    member = _require_member(archive, basename)
    payload = json.loads(archive.read(member).decode("utf-8-sig"))
    rows = payload.get("row")
    if not isinstance(rows, list):
        raise ValueError(f"O*NET JSON member '{basename}' does not contain a row list")
    return [row for row in rows if isinstance(row, dict)]


def _detect_format(zip_path: Path) -> str:
    with zipfile.ZipFile(zip_path) as archive:
        if _member_by_basename(archive, _FULLDB_JSON_OCCUPATIONS_MEMBER) and _member_by_basename(
            archive, _FULLDB_JSON_JOB_TITLES_MEMBER
        ):
            return "fulldb_json"
        if _member_by_basename(archive, _FULLDB_TEXT_OCCUPATIONS_MEMBER) and (
            _member_by_basename(archive, _FULLDB_TEXT_JOB_TITLES_MEMBER)
            or _member_by_basename(archive, _FULLDB_TEXT_ALT_TITLES_MEMBER)
        ):
            return "fulldb_text"
        if _member_by_basename(archive, _LISTINGS_OCCUPATIONS_MEMBER):
            return "listings"
    raise ValueError(
        "Unrecognised O*NET zip format. Expected the current full database "
        "(JSON or text) or legacy OccupationalListings.zip."
    )


def _read_listings_format(zip_path: Path) -> tuple[list[dict], list[dict], list[dict]]:
    occupation_rows = _read_xlsx_from_zip(zip_path, _LISTINGS_OCCUPATIONS_MEMBER)
    title_rows = _read_xlsx_from_zip(zip_path, _LISTINGS_ALT_TITLES_MEMBER)
    occupations = [
        {
            "code": str(row["O*NET-SOC 2019 Code"]).strip(),
            "title": str(row["O*NET-SOC 2019 Title"]).strip(),
            "description": str(row.get("O*NET-SOC 2019 Description") or "").strip(),
        }
        for row in occupation_rows
    ]
    job_titles = [
        {
            "occupation_code": str(row["O*NET-SOC 2019 Code"]).strip(),
            "job_title": str(row["Alternate Title"]).strip(),
            "short_title": "",
            "sources": "",
        }
        for row in title_rows
    ]
    return occupations, job_titles, []


def _read_fulldb_text_format(zip_path: Path) -> tuple[list[dict], list[dict], list[dict]]:
    with zipfile.ZipFile(zip_path) as archive:
        occupation_rows = _read_txt_table(archive, _FULLDB_TEXT_OCCUPATIONS_MEMBER)
        job_title_member = (
            _FULLDB_TEXT_JOB_TITLES_MEMBER
            if _member_by_basename(archive, _FULLDB_TEXT_JOB_TITLES_MEMBER)
            else _FULLDB_TEXT_ALT_TITLES_MEMBER
        )
        title_rows = _read_txt_table(archive, job_title_member)
        reported_rows = (
            _read_txt_table(archive, _FULLDB_TEXT_REPORTED_TITLES_MEMBER)
            if _member_by_basename(archive, _FULLDB_TEXT_REPORTED_TITLES_MEMBER)
            else []
        )
    occupations = [
        {
            "code": str(row.get("O*NET-SOC Code") or "").strip(),
            "title": str(row.get("Title") or "").strip(),
            "description": str(row.get("Description") or "").strip(),
        }
        for row in occupation_rows
    ]
    title_column = (
        "Job Title" if job_title_member == _FULLDB_TEXT_JOB_TITLES_MEMBER else "Alternate Title"
    )
    job_titles = [
        {
            "occupation_code": str(row.get("O*NET-SOC Code") or "").strip(),
            "job_title": str(row.get(title_column) or "").strip(),
            "short_title": _usable_title(row.get("Short Title")),
            "sources": str(row.get("Source(s)") or "").strip(),
        }
        for row in title_rows
    ]
    reported_titles = [
        {
            "occupation_code": str(row.get("O*NET-SOC Code") or "").strip(),
            "reported_job_title": str(row.get("Reported Job Title") or "").strip(),
            "shown_in_my_next_move": str(row.get("Shown in My Next Move") or "").strip(),
        }
        for row in reported_rows
    ]
    return occupations, job_titles, reported_titles


def _read_fulldb_json_format(zip_path: Path) -> tuple[list[dict], list[dict], list[dict]]:
    with zipfile.ZipFile(zip_path) as archive:
        occupation_rows = _read_json_table(archive, _FULLDB_JSON_OCCUPATIONS_MEMBER)
        title_rows = _read_json_table(archive, _FULLDB_JSON_JOB_TITLES_MEMBER)
        reported_rows = (
            _read_json_table(archive, _FULLDB_JSON_REPORTED_TITLES_MEMBER)
            if _member_by_basename(archive, _FULLDB_JSON_REPORTED_TITLES_MEMBER)
            else []
        )
    occupations = [
        {
            "code": str(row.get("onetsoc_code") or "").strip(),
            "title": str(row.get("title") or "").strip(),
            "description": str(row.get("description") or "").strip(),
        }
        for row in occupation_rows
    ]
    job_titles = [
        {
            "occupation_code": str(row.get("onetsoc_code") or "").strip(),
            "job_title": str(row.get("job_title") or "").strip(),
            "short_title": _usable_title(row.get("short_title")),
            "sources": str(row.get("sources") or "").strip(),
        }
        for row in title_rows
    ]
    reported_titles = [
        {
            "occupation_code": str(row.get("onetsoc_code") or "").strip(),
            "reported_job_title": str(row.get("reported_job_title") or "").strip(),
            "shown_in_my_next_move": str(row.get("shown_in_my_next_move") or "").strip(),
        }
        for row in reported_rows
    ]
    return occupations, job_titles, reported_titles


def _dataset_fingerprint(by_normalized_title: dict[str, list[dict[str, Any]]]) -> str:
    canonical = json.dumps(
        by_normalized_title, ensure_ascii=False, sort_keys=True, separators=(",", ":")
    )
    return hashlib.sha256(canonical.encode("utf-8")).hexdigest()


def build_taxonomy(
    zip_path: Path,
    *,
    database_release: str | None = None,
    source_url: str | None = None,
) -> dict[str, Any]:
    """Return normalized occupations/title index plus source metadata."""
    fmt = _detect_format(zip_path)
    if fmt == "fulldb_json":
        raw_occupations, raw_job_titles, reported_titles = _read_fulldb_json_format(zip_path)
        source_label = "O*NET Database JSON"
    elif fmt == "fulldb_text":
        raw_occupations, raw_job_titles, reported_titles = _read_fulldb_text_format(zip_path)
        source_label = "O*NET Database text"
    else:
        raw_occupations, raw_job_titles, reported_titles = _read_listings_format(zip_path)
        source_label = "O*NET OccupationalListings.zip"

    release = database_release or _infer_database_release(zip_path)
    occupations = [
        {
            "code": row["code"],
            "title": row["title"],
            "normalized_title": normalize_title(row["title"]),
            "description": row["description"],
        }
        for row in raw_occupations
        if row.get("code") and row.get("title")
    ]
    occupation_title_by_code = {row["code"]: row["title"] for row in occupations}
    preferred_pairs = {
        (normalize_title(row.get("reported_job_title")), row.get("occupation_code"))
        for row in reported_titles
        if str(row.get("shown_in_my_next_move") or "").upper() == "Y"
        and normalize_title(row.get("reported_job_title"))
        and row.get("occupation_code")
    }

    by_normalized_title: dict[str, list[dict[str, Any]]] = {}
    for occupation in occupations:
        by_normalized_title.setdefault(occupation["normalized_title"], []).append(
            {
                "occupation_code": occupation["code"],
                "occupation_title": occupation["title"],
                "source": "occupation_title",
            }
        )

    title_variant_count = 0
    for row in raw_job_titles:
        code = str(row.get("occupation_code") or "").strip()
        occupation_title = occupation_title_by_code.get(code, "")
        if not code or not occupation_title:
            continue
        for variant in (row.get("job_title"), row.get("short_title")):
            display_title = _usable_title(variant)
            normalized = normalize_title(display_title)
            if not normalized:
                continue
            match: dict[str, Any] = {
                "occupation_code": code,
                "occupation_title": occupation_title,
                "matched_title": display_title,
                "source": "job_title",
            }
            if (normalized, code) in preferred_pairs:
                match["target_query_preferred"] = True
            bucket = by_normalized_title.setdefault(normalized, [])
            dedupe_key = (code, display_title.casefold(), "job_title")
            existing_keys = {
                (
                    str(item.get("occupation_code") or ""),
                    str(item.get("matched_title") or item.get("occupation_title") or "").casefold(),
                    str(item.get("source") or ""),
                )
                for item in bucket
            }
            if dedupe_key not in existing_keys:
                bucket.append(match)
                title_variant_count += 1

    ordered_index = {key: by_normalized_title[key] for key in sorted(by_normalized_title)}
    fingerprint = _dataset_fingerprint(ordered_index)
    metadata = {
        "source": source_label,
        "source_url": source_url or "",
        "taxonomy_version": TAXONOMY_VERSION,
        "database_release": release or "",
        "dataset_fingerprint": fingerprint,
        "format": fmt,
        "record_counts": {
            "occupations": len(occupations),
            "job_title_rows": len(raw_job_titles),
            "title_variants": title_variant_count,
            "normalized_title_keys": len(ordered_index),
            "target_preferred_pairs": len(preferred_pairs),
        },
    }
    return {
        "metadata": metadata,
        "occupations": occupations,
        "by_normalized_title": ordered_index,
    }


def validate_taxonomy(taxonomy: dict[str, Any]) -> None:
    """Reject obviously partial/corrupt full-database imports before replacement."""
    metadata = taxonomy.get("metadata") or {}
    counts = metadata.get("record_counts") or {}
    if not taxonomy.get("by_normalized_title") or not taxonomy.get("occupations"):
        raise ValueError("O*NET taxonomy build produced no occupations/title index")
    if str(metadata.get("format") or "").startswith("fulldb_"):
        if int(counts.get("occupations") or 0) < 900:
            raise ValueError("O*NET full database contains unexpectedly few occupations")
        if int(counts.get("job_title_rows") or 0) < 10_000:
            raise ValueError("O*NET full database contains unexpectedly few job-title rows")
        if not metadata.get("database_release"):
            raise ValueError("O*NET full database release could not be identified")
    if len(str(metadata.get("dataset_fingerprint") or "")) != 64:
        raise ValueError("O*NET taxonomy dataset fingerprint is missing or invalid")


def write_taxonomy_files(
    zip_path: Path,
    output_dir: Path = DEFAULT_OUTPUT_DIR,
    *,
    database_release: str | None = None,
    source_url: str | None = None,
) -> dict[str, Any]:
    """Build, validate, then atomically replace generated O*NET reference files."""
    taxonomy = build_taxonomy(zip_path, database_release=database_release, source_url=source_url)
    validate_taxonomy(taxonomy)
    output_dir.mkdir(parents=True, exist_ok=True)
    files = {
        "onet_occupations.json": {
            "metadata": taxonomy["metadata"],
            "occupations": taxonomy["occupations"],
        },
        "onet_index.json": {
            "metadata": taxonomy["metadata"],
            "normalization": "lowercase; ampersand to and; non-alphanumeric collapsed to spaces",
            "by_normalized_title": taxonomy["by_normalized_title"],
        },
    }
    with TemporaryDirectory(dir=output_dir.parent) as tmp_dir:
        tmp_root = Path(tmp_dir)
        for filename, payload in files.items():
            (tmp_root / filename).write_text(
                json.dumps(payload, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
            )
        for filename in files:
            (tmp_root / filename).replace(output_dir / filename)
    obsolete = output_dir / "onet_alternate_titles.json"
    if obsolete.exists():
        obsolete.unlink()
    counts = taxonomy["metadata"]["record_counts"]
    print(
        f"Generated O*NET taxonomy in {output_dir}\n"
        f"  Source:           {taxonomy['metadata']['source']}\n"
        f"  Database release: {taxonomy['metadata']['database_release'] or 'not supplied'}\n"
        f"  Occupations:      {counts['occupations']}\n"
        f"  Job-title rows:   {counts['job_title_rows']}\n"
        f"  Index keys:       {counts['normalized_title_keys']}"
    )
    return taxonomy["metadata"]


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("zip_path", type=Path, help="Path to an O*NET source zip")
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--database-release", default=None)
    parser.add_argument("--source-url", default=None)
    args = parser.parse_args()
    write_taxonomy_files(
        args.zip_path,
        args.output_dir,
        database_release=args.database_release,
        source_url=args.source_url,
    )


if __name__ == "__main__":
    main()
